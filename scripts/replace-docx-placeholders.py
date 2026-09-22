#!/usr/bin/env python3
"""Replace simple DOCX placeholders in paragraphs and table cells.

This preserves the first run's formatting for a paragraph. For complex templates,
use a manifest and test the rendered result; Word may split placeholders across runs.
"""
import argparse
from pathlib import Path
from docx import Document


def replace_paragraph(paragraph, values):
    original = paragraph.text
    updated = original
    for key, value in values.items():
        updated = updated.replace("{{" + key + "}}", str(value))
    if updated == original:
        return False
    if paragraph.runs:
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(updated)
    return True


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("template")
    parser.add_argument("output")
    parser.add_argument("--value", action="append", default=[], metavar="KEY=VALUE")
    args = parser.parse_args()
    values = dict(item.split("=", 1) for item in args.value if "=" in item)
    doc = Document(args.template)
    count = 0
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        count += replace_paragraph(paragraph, values)
    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(f"replaced_paragraphs: {count}")
    print(f"output: {args.output}")


if __name__ == "__main__":
    main()
