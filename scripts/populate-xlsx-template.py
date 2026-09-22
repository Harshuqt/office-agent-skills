#!/usr/bin/env python3
"""Populate explicit worksheet cells from a JSON object.

Mapping format: {"Sheet1!B2": "value", "Inputs!B3": 123}
"""
import argparse
import json
from pathlib import Path
from openpyxl import load_workbook


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("template")
    parser.add_argument("mapping_json")
    parser.add_argument("output")
    args = parser.parse_args()
    mapping = json.loads(Path(args.mapping_json).read_text(encoding="utf-8"))
    keep_vba = Path(args.template).suffix.lower() == ".xlsm"
    workbook = load_workbook(args.template, keep_vba=keep_vba)
    for address, value in mapping.items():
        if "!" not in address:
            raise SystemExit(f"Mapping must use Sheet!Cell notation: {address}")
        sheet_name, cell = address.rsplit("!", 1)
        if sheet_name not in workbook.sheetnames:
            raise SystemExit(f"Unknown sheet: {sheet_name}")
        workbook[sheet_name][cell] = value
    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(f"updated_cells: {len(mapping)}")
    print(f"output: {args.output}")


if __name__ == "__main__":
    main()
